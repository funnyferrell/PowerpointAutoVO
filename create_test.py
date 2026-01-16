#!/usr/bin/env python3
"""
Create a test PPTX with animations for testing the converter.

This creates a presentation with:
- Slide 1: Title (no animations)
- Slide 2: 3 boxes that appear on click (with [CLICK] markers in notes)
- Slide 3: Bullet points that appear on click
- Slide 4: Thank you (no notes)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

# OOXML namespace
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def add_click_animations(slide, shape_id_groups):
    """
    Add click-triggered appear animations to shape groups.
    Each group of shape IDs will appear together on a single click.
    
    shape_id_groups: list of lists, e.g., [[box1_id, text1_id], [box2_id, text2_id], ...]
    """
    # Build animation XML
    child_tn_content = []
    node_id = 3
    bld_p_content = []
    
    for group in shape_id_groups:
        # All shapes in this group animate together on one click
        group_animations = []
        
        for idx, shape_id in enumerate(group):
            # First shape in group is OnClick, rest are WithPrevious
            trigger_delay = "indefinite" if idx == 0 else "0"
            
            anim = f'''
                                <p:par>
                                    <p:cTn id="{node_id+2}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{'clickEffect' if idx == 0 else 'withEffect'}">
                                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                        <p:childTnLst>
                                            <p:set>
                                                <p:cBhvr>
                                                    <p:cTn id="{node_id+3}" dur="1" fill="hold">
                                                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                                    </p:cTn>
                                                    <p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl>
                                                    <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                                                </p:cBhvr>
                                                <p:to><p:strVal val="visible"/></p:to>
                                            </p:set>
                                        </p:childTnLst>
                                    </p:cTn>
                                </p:par>'''
            group_animations.append(anim)
            bld_p_content.append(f'<p:bldP xmlns:p="{P_NS}" spid="{shape_id}" grpId="0"/>')
            node_id += 4
        
        # Wrap all animations in this group in a single click container
        child_tn_content.append(f'''
        <p:par xmlns:p="{P_NS}">
            <p:cTn id="{node_id}" fill="hold">
                <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
                <p:childTnLst>
                    <p:par>
                        <p:cTn id="{node_id+1}" fill="hold">
                            <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                            <p:childTnLst>
                                {"".join(group_animations)}
                            </p:childTnLst>
                        </p:cTn>
                    </p:par>
                </p:childTnLst>
            </p:cTn>
        </p:par>
        ''')
        node_id += 2
    
    timing_xml = f'''
    <p:timing xmlns:p="{P_NS}">
        <p:tnLst>
            <p:par>
                <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
                    <p:childTnLst>
                        <p:seq concurrent="1" nextAc="seek">
                            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                                <p:childTnLst>
                                    {"".join(child_tn_content)}
                                </p:childTnLst>
                            </p:cTn>
                            <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
                            <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
                        </p:seq>
                    </p:childTnLst>
                </p:cTn>
            </p:par>
        </p:tnLst>
        <p:bldLst>
            {"".join(bld_p_content)}
        </p:bldLst>
    </p:timing>
    '''
    
    timing_elem = etree.fromstring(timing_xml)
    
    # Remove existing timing
    existing = slide._element.find(f'.//{{{P_NS}}}timing')
    if existing is not None:
        existing.getparent().remove(existing)
    
    slide._element.append(timing_elem)


# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# =============================================================================
# Slide 1: Title
# =============================================================================
slide1 = prs.slides.add_slide(blank_layout)

title = slide1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
tf = title.text_frame
p = tf.paragraphs[0]
p.text = "Animation Demo"
p.font.size = Pt(54)
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

subtitle = slide1.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(1))
tf = subtitle.text_frame
p = tf.paragraphs[0]
p.text = "Testing PowerPoint Animation Capture"
p.font.size = Pt(24)
p.alignment = PP_ALIGN.CENTER
p.font.color.rgb = RGBColor(100, 100, 100)

notes1 = slide1.notes_slide
notes1.notes_text_frame.text = "Welcome to this animation demo. This presentation will test our ability to capture click-triggered animations."

# =============================================================================
# Slide 2: Three colored boxes appearing on click
# =============================================================================
slide2 = prs.slides.add_slide(blank_layout)

# Title (always visible)
title2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
tf = title2.text_frame
p = tf.paragraphs[0]
p.text = "Three Steps"
p.font.size = Pt(44)
p.font.bold = True

# Box 1 - Red
box1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(3.5), Inches(3))
box1.fill.solid()
box1.fill.fore_color.rgb = RGBColor(220, 53, 69)
box1.line.fill.background()

text1 = slide2.shapes.add_textbox(Inches(1.3), Inches(3), Inches(3), Inches(1))
tf = text1.text_frame
p = tf.paragraphs[0]
p.text = "Step 1"
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)
p.alignment = PP_ALIGN.CENTER

# Box 2 - Green
box2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(2), Inches(3.5), Inches(3))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(40, 167, 69)
box2.line.fill.background()

text2 = slide2.shapes.add_textbox(Inches(5.3), Inches(3), Inches(3), Inches(1))
tf = text2.text_frame
p = tf.paragraphs[0]
p.text = "Step 2"
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)
p.alignment = PP_ALIGN.CENTER

# Box 3 - Blue
box3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9), Inches(2), Inches(3.5), Inches(3))
box3.fill.solid()
box3.fill.fore_color.rgb = RGBColor(0, 123, 255)
box3.line.fill.background()

text3 = slide2.shapes.add_textbox(Inches(9.3), Inches(3), Inches(3), Inches(1))
tf = text3.text_frame
p = tf.paragraphs[0]
p.text = "Step 3"
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)
p.alignment = PP_ALIGN.CENTER

# Add animations - each box+text appears together on one click (3 clicks total)
add_click_animations(slide2, [
    [box1.shape_id, text1.shape_id],  # Click 1: Step 1 (red)
    [box2.shape_id, text2.shape_id],  # Click 2: Step 2 (green)
    [box3.shape_id, text3.shape_id],  # Click 3: Step 3 (blue)
])

notes2 = slide2.notes_slide
notes2.notes_text_frame.text = """Let me show you our three-step process. [CLICK] Step one is preparation and planning. [CLICK] Step two is execution and implementation. [CLICK] Step three is review and optimization."""

# =============================================================================
# Slide 3: Bullet points appearing on click
# =============================================================================
slide3 = prs.slides.add_slide(blank_layout)

title3 = slide3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
tf = title3.text_frame
p = tf.paragraphs[0]
p.text = "Key Takeaways"
p.font.size = Pt(44)
p.font.bold = True

# Bullet 1 (always visible)
pt1 = slide3.shapes.add_textbox(Inches(1), Inches(1.5), Inches(10), Inches(0.8))
tf = pt1.text_frame
p = tf.paragraphs[0]
p.text = "• Start with clear objectives"
p.font.size = Pt(28)

# Bullet 2 (click to reveal)
pt2 = slide3.shapes.add_textbox(Inches(1), Inches(2.4), Inches(10), Inches(0.8))
tf = pt2.text_frame
p = tf.paragraphs[0]
p.text = "• Measure progress regularly"
p.font.size = Pt(28)
p.font.color.rgb = RGBColor(0, 100, 0)

# Bullet 3 (click to reveal)
pt3 = slide3.shapes.add_textbox(Inches(1), Inches(3.3), Inches(10), Inches(0.8))
tf = pt3.text_frame
p = tf.paragraphs[0]
p.text = "• Iterate based on feedback"
p.font.size = Pt(28)
p.font.color.rgb = RGBColor(0, 100, 0)

# Each bullet is a separate click (2 clicks total)
add_click_animations(slide3, [
    [pt2.shape_id],  # Click 1: Second bullet
    [pt3.shape_id],  # Click 2: Third bullet
])

notes3 = slide3.notes_slide
notes3.notes_text_frame.text = """Here are the key takeaways. First, always start with clear objectives. [CLICK] Second, measure your progress regularly to stay on track. [CLICK] And third, iterate based on the feedback you receive."""

# =============================================================================
# Slide 4: Thank you
# =============================================================================
slide4 = prs.slides.add_slide(blank_layout)

thanks = slide4.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(1.5))
tf = thanks.text_frame
p = tf.paragraphs[0]
p.text = "Thank You!"
p.font.size = Pt(60)
p.font.bold = True
p.alignment = PP_ALIGN.CENTER

# No notes = silent slide

# =============================================================================
# Save
# =============================================================================
prs.save("test_presentation.pptx")

print("Created: test_presentation.pptx")
print()
print("Slide structure:")
print("  1. Title slide (no animations)")
print("  2. Three Steps - 3 [CLICK] markers, 6 animated shapes")
print("  3. Key Takeaways - 2 [CLICK] markers, 2 animated shapes")
print("  4. Thank You (no notes)")
print()
print("To test:")
print("  python convert.py test_presentation.pptx output.mp4 --provider test")
